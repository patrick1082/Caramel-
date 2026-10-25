from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from predictor.cartelera import Analisis, FilaComparacion, Oportunidad
from predictor.estadisticas import calcular
from predictor.evaluacion import TOTAL, MetricasFuente, ResultadoValor, cruzar, por_liga
from predictor.exportar import (
    COLUMNAS_APUESTAS,
    COLUMNAS_MOVIMIENTOS,
    COLUMNAS_PREDICCIONES,
    exportar_excel,
    resumen_contraste,
    tabla_backtest_valor,
    tabla_cara_a_cara,
    tabla_contraste,
    tabla_cuotas,
    tabla_estadisticas,
    tabla_evaluacion,
    tabla_fuentes,
    tabla_jugadores,
    tabla_predicciones,
    tabla_valor_del_dia,
)
from predictor.ligas import ZONA_CHILE
from predictor.modelos import (
    Deporte,
    Estado,
    Fuente,
    Jugador,
    Mercado,
    Movimiento,
    Partido,
    Prediccion,
    PrediccionPropia,
    ProbabilidadLinea,
    Pronostico3,
    ResultadoHistorico,
    Rol,
    TipoMovimiento,
    Valor,
)
from predictor.valor import valorar


def _abrir(bitacora):
    contenido = exportar_excel(bitacora.apuestas(), bitacora.movimientos(), bitacora.resumen())
    return load_workbook(BytesIO(contenido))


def _filas(hoja):
    return list(hoja.iter_rows(values_only=True))


def test_exporta_cinco_hojas(bitacora):
    libro = _abrir(bitacora)

    assert libro.sheetnames == ["Apuestas", "Movimientos", "Resumen", "Predicciones", "Contraste"]
    assert _filas(libro["Predicciones"]) == [COLUMNAS_PREDICCIONES]
    assert _filas(libro["Contraste"]) == [("Fuente", "Partidos", "Acierto (%)", "Brier")]


def test_exporta_valor_de_las_apuestas_y_contraste(bitacora, crear_apuesta):
    apuesta = bitacora.registrar_apuesta(crear_apuesta(prob_modelo=0.6, valor=Valor.CONFIRMADO))
    bitacora.liquidar(apuesta.id, Estado.GANADA)
    fuentes = {"Modelo": MetricasFuente(partidos=2, aciertos=1, brier_suma=1.2)}

    libro = load_workbook(
        BytesIO(exportar_excel(bitacora.apuestas(), [], bitacora.resumen(), fuentes=fuentes))
    )

    apuestas = _filas(libro["Apuestas"])
    fila = dict(zip(apuestas[0], apuestas[1], strict=True))
    assert fila["Modelo (%)"] == 60.0 and fila["Valor"] == "✅ Confirmado"
    assert any(f[0] == "Valor: ✅ Confirmado" for f in _filas(libro["Resumen"]))
    assert _filas(libro["Contraste"])[1] == ("Modelo", 2, 50.0, 0.6)


def test_textos_con_formula_se_exportan_como_texto(bitacora, crear_apuesta):
    bitacora.registrar_apuesta(crear_apuesta(nota='=HYPERLINK("http://x","clic")'))
    hoja = _abrir(bitacora)["Apuestas"]

    celdas = [c for fila in hoja.iter_rows() for c in fila if str(c.value).startswith("=")]
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_exporta_vacio_con_encabezados(bitacora):
    libro = _abrir(bitacora)

    assert _filas(libro["Apuestas"]) == [tuple(COLUMNAS_APUESTAS.values())]
    assert _filas(libro["Movimientos"]) == [tuple(COLUMNAS_MOVIMIENTOS.values())]


def test_exporta_apuestas_y_movimientos(bitacora, crear_apuesta):
    apuesta = bitacora.registrar_apuesta(crear_apuesta(monto=1000))
    bitacora.registrar_apuesta(crear_apuesta(monto=500, colocada=False))
    bitacora.liquidar(apuesta.id, Estado.GANADA)
    bitacora.registrar_movimiento(Movimiento(tipo=TipoMovimiento.DEPOSITO, monto=3_000))

    libro = _abrir(bitacora)
    apuestas = _filas(libro["Apuestas"])
    encabezados = apuestas[0]
    primera = dict(zip(encabezados, apuestas[1], strict=True))

    assert len(apuestas) == 3
    assert primera["Partido"] == "Lakers vs Celtics"
    assert primera["Cuota"] == 1.85
    assert primera["Monto (CLP)"] == 1000
    assert primera["Estado"] == "ganada"
    assert primera["Resultado neto (CLP)"] == 850
    assert primera["Colocada"] == "Sí"
    assert len(_filas(libro["Movimientos"])) == 2


def test_resumen_incluye_saldo_y_desglose(bitacora, crear_apuesta):
    apuesta = bitacora.registrar_apuesta(crear_apuesta(monto=1000))
    bitacora.liquidar(apuesta.id, Estado.PERDIDA)

    filas = _filas(_abrir(bitacora)["Resumen"])
    valores = {fila[0]: fila[1] for fila in filas if fila[0]}

    assert valores["Saldo actual (CLP)"] == 9_000
    assert valores["ROI (%)"] == -100
    assert any(fila[0] == "NBA" for fila in filas)


def _fila(mercado, seleccion, linea=None, betano="1.9"):
    return FilaComparacion(
        mercado=mercado,
        seleccion=seleccion,
        linea=Decimal(linea) if linea else None,
        betano=Decimal(betano) if betano else None,
        pinnacle=Decimal("2.0"),
        mejor=Decimal("2.1"),
        casa_mejor="Betfair",
        promedio=Decimal("1.95"),
        casas=5,
    )


def test_tabla_cuotas_traduce_empate_y_formatea_lineas():
    tabla = tabla_cuotas(
        [
            _fila(Mercado.GANADOR, "Draw"),
            _fila(Mercado.TOTALES, "Over", "2.5", betano=None),
            _fila(Mercado.HANDICAP, "Arsenal", "0.75"),
        ]
    )

    assert list(tabla["Mercado"]) == ["Ganador", "Over/Under", "Hándicap"]
    assert list(tabla["Selección"]) == ["Empate", "Over", "Arsenal"]
    assert list(tabla["Línea"]) == ["", "2.5", "+0.75"]
    assert tabla["Betano"].isna().tolist() == [False, True, False]
    assert tabla.loc[0, "Mejor"] == 2.1


def test_tabla_estadisticas_lado_a_lado():
    historial = [
        ResultadoHistorico(
            "Premier League",
            "2627",
            date(2026, 9, 1),
            "Arsenal",
            "Everton",
            2,
            1,
            {"tiros": (12.0, 5.0)},
        )
    ]
    analisis = Analisis(calcular("Arsenal", historial), None, historial)

    tabla = tabla_estadisticas(analisis, "Arsenal", "Leeds United")

    assert list(tabla.columns) == ["Concepto", "Arsenal", "Leeds United"]
    fila = dict(zip(tabla["Concepto"], tabla["Arsenal"], strict=True))
    assert fila["G-E-P"] == "1-0-0"
    assert fila["Anota / recibe"] == "2 / 1 (1 PJ)"
    assert fila["Como local · como visita"] == "2 / 1 (1 PJ) · —"
    assert fila["Tiros"] == "12 / 5"
    assert tabla["Leeds United"].iloc[0] == "Sin datos"
    assert list(tabla_cara_a_cara(historial).iloc[0]) == [
        "2026-09-01",
        "Arsenal",
        "2 - 1",
        "Everton",
        "2627",
    ]


INICIO = datetime(2026, 10, 20, 16, 0, tzinfo=ZONA_CHILE)


def _prediccion(partido_id="p1", local="Arsenal", visitante="Everton"):
    return Prediccion(
        partido_id=partido_id,
        liga="Premier League",
        local=local,
        visitante=visitante,
        inicio=INICIO,
        modelo="poisson-1",
        creada=INICIO,
        esperado_local=1.8,
        esperado_visitante=0.9,
        prob_local=0.5,
        prob_empate=0.25,
        prob_visitante=0.25,
        marcador_probable="1-0",
        lineas=[
            ProbabilidadLinea(Mercado.TOTALES, "Over", 2.5, 0.48),
            ProbabilidadLinea(Mercado.TOTALES, "Under", 2.5, 0.52),
        ],
        linea_total=2.5,
    )


def test_tabla_cuotas_con_probabilidad_del_modelo_y_cuota_justa():
    tabla = tabla_cuotas(
        [
            _fila(Mercado.GANADOR, "Arsenal"),
            _fila(Mercado.GANADOR, "Draw"),
            _fila(Mercado.TOTALES, "Under", "2.5"),
            _fila(Mercado.HANDICAP, "Arsenal", "-1.5"),  # el modelo no la calculó
        ],
        _prediccion(),
    )

    assert tabla["Modelo (%)"].tolist()[:3] == [50.0, 25.0, 52.0]
    assert tabla["Cuota justa"].tolist()[:3] == [2.0, 4.0, 1.92]
    assert tabla[["Modelo (%)", "Cuota justa"]].iloc[3].isna().all()
    assert "Modelo (%)" not in tabla_cuotas([_fila(Mercado.GANADOR, "Arsenal")]).columns


def test_tabla_jugadores_de_ambos_equipos():
    tabla = tabla_jugadores(
        "Arsenal",
        [Jugador("Premier League", "Arsenal", "Saka", Rol.DESTACADO, "5 goles · 8 PJ", 5)],
        "Everton",
        [Jugador("Premier League", "Everton", "Pickford", Rol.LESIONADO, "Out · Knee")],
    )

    assert tabla.values.tolist() == [
        ["Arsenal", "Destacado", "Saka", "5 goles · 8 PJ"],
        ["Everton", "Lesionado", "Pickford", "Out · Knee"],
    ]


def test_tabla_evaluacion_y_predicciones():
    historial = [
        ResultadoHistorico("Premier League", "2627", INICIO.date(), "Arsenal", "Everton", 2, 1)
    ]
    pendiente = _prediccion("p2", "Chelsea", "Fulham")
    evaluadas = cruzar([_prediccion(), pendiente], historial)

    evaluacion = tabla_evaluacion(por_liga(evaluadas))
    predicciones = tabla_predicciones([_prediccion(), pendiente], evaluadas)

    assert evaluacion["Liga"].tolist() == ["Premier League", TOTAL]
    liga = evaluacion.iloc[0]
    assert liga["Partidos"] == 1 and liga["Acierto modelo (%)"] == 100.0
    assert liga["Brier modelo"] == round(0.5**2 + 0.25**2 + 0.25**2, 3)
    assert liga["Acierto over/under (%)"] == 0.0  # 3 goles, pero el modelo daba Under 52 %
    assert evaluacion["Error medio"].isna().tolist() == [False, True]
    primera = dict(zip(COLUMNAS_PREDICCIONES, predicciones.iloc[0], strict=True))
    assert primera["Inicio"] == "2026-10-20 16:00"
    assert primera["Esperado"] == "1.8 - 0.9" and primera["Over (%)"] == 48.0
    assert primera["Resultado"] == "2 - 1" and primera["Acierto"] == "Sí"
    assert predicciones.iloc[1]["Resultado"] == "" and predicciones.iloc[1]["Acierto"] == ""


def _ganador_1x2(seleccion, betano, pinnacle):
    return FilaComparacion(
        mercado=Mercado.GANADOR,
        seleccion=seleccion,
        linea=None,
        betano=Decimal(betano),
        pinnacle=Decimal(pinnacle),
        mejor=Decimal(betano),
        casa_mejor="Betano",
        promedio=Decimal(pinnacle),
        casas=3,
    )


FILAS_1X2 = [
    _ganador_1x2("Arsenal", "2.30", "2.10"),
    _ganador_1x2("Draw", "3.40", "3.50"),
    _ganador_1x2("Everton", "3.60", "3.80"),
]


def test_tabla_cuotas_con_valoraciones():
    prediccion = _prediccion()
    tabla = tabla_cuotas(FILAS_1X2, prediccion, valorar(FILAS_1X2, prediccion, "Arsenal"))

    assert tabla["Valor"].tolist()[0] == "✅ Confirmado"  # 0,5 × 2,30 = EV +15 %
    assert tabla["EV (%)"].tolist()[0] == 15.0
    assert tabla["Mercado (%)"].notna().all()
    assert tabla["Valor"].tolist()[1:] == ["—", "—"]


def test_tabla_valor_del_dia():
    partido = Partido("p1", Deporte.FUTBOL, "Premier League", "Arsenal", "Everton", INICIO)
    valoracion = valorar(FILAS_1X2, _prediccion(), "Arsenal")[0]

    tabla = tabla_valor_del_dia([Oportunidad(partido, valoracion)])

    fila = tabla.iloc[0]
    assert (fila["Hora"], fila["Partido"], fila["Selección"]) == (
        "16:00",
        "Arsenal vs Everton",
        "Arsenal",
    )
    assert (fila["Betano"], fila["EV (%)"], fila["Valor"]) == (2.3, 15.0, "✅ Confirmado")


def _pronostico(fuente, local, empate, visita):
    return Pronostico3("p1", fuente, local, empate, visita, INICIO)


def _propia(ganador, marcador=(None, None)):
    return PrediccionPropia(
        "p1", "Premier League", "Arsenal", "Everton", INICIO, ganador, INICIO, *marcador
    )


def test_tabla_contraste_lado_a_lado():
    pronosticos = {
        Fuente.MERCADO: _pronostico(Fuente.MERCADO, 0.45, 0.28, 0.27),
        Fuente.MODELO: _pronostico(Fuente.MODELO, 0.5, 0.25, 0.25),
    }

    tabla = tabla_contraste(
        "Arsenal", "Everton", _propia("visitante", (0, 1)), pronosticos, _prediccion()
    )

    assert tabla["Opinión"].tolist() == ["Mía", "Modelo", "Mercado"]
    assert tabla["Favorito"].tolist() == ["Everton", "Arsenal", "Arsenal"]
    assert tabla["Marcador"].tolist() == ["0 - 1", "1-0", ""]
    assert tabla.loc[1, "Empate (%)"] == 25.0 and tabla[["Local (%)"]].iloc[0].isna().all()


def test_tabla_contraste_nba_sin_columna_de_empate():
    tabla = tabla_contraste(
        "Lakers", "Celtics", None, {Fuente.ESPN: _pronostico(Fuente.ESPN, 0.4, None, 0.6)}
    )

    assert list(tabla.columns) == ["Opinión", "Local (%)", "Visita (%)", "Favorito", "Marcador"]
    assert tabla.iloc[0]["Favorito"] == "Celtics"


def test_resumen_contraste():
    modelo = _pronostico(Fuente.MODELO, 0.5, 0.25, 0.25)
    mercado = _pronostico(Fuente.MERCADO, 0.3, 0.4, 0.3)
    local = _propia("local")

    assert resumen_contraste("Arsenal", "Everton", local, {Fuente.MODELO: modelo}) == (
        "Coinciden tú y el modelo: gana Arsenal"
    )
    assert (
        resumen_contraste(
            "Arsenal", "Everton", local, {Fuente.MODELO: modelo, Fuente.MERCADO: mercado}
        )
        == "Gana Arsenal según tú y el modelo · Empate según el mercado"
    )
    assert resumen_contraste("Arsenal", "Everton", None, {Fuente.MODELO: modelo}) == ""


def test_tabla_fuentes_y_backtest_valor():
    fuentes = tabla_fuentes(
        {"Mía": MetricasFuente(partidos=4, aciertos=3, brier_suma=None), "Modelo": MetricasFuente()}
    )
    backtest = tabla_backtest_valor(
        {
            TOTAL: ResultadoValor(
                partidos=10, apuestas=4, aciertos=1, ganancia=-1.5, ganancia_favorito=0.5
            )
        }
    )

    assert fuentes.iloc[0].tolist()[:3] == ["Mía", 4, 75.0] and fuentes["Brier"].isna().all()
    assert backtest.iloc[0].tolist() == [TOTAL, 10, 4, 25.0, -37.5, 5.0]
